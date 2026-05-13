import asyncio
import os
import sys
from datetime import datetime
from azure.identity.aio import DefaultAzureCredential
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Add project root to sys.path to import backend modules
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from backend.settings import app_settings
from backend.history.cosmosdbservice import CosmosConversationClient

COLUMNS = [
    ("User ID",              28),
    ("Conversation ID",      36),
    ("Conversation Title",   30),
    ("Conversation Created", 22),
    ("Message #",           10),
    ("Role",                12),
    ("Timestamp",           22),
    ("Message Content",     80),
]

HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")
USER_FILL    = PatternFill("solid", fgColor="D6E4F0")
AI_FILL      = PatternFill("solid", fgColor="E8F5E9")
TOOL_FILL    = PatternFill("solid", fgColor="FFF9C4")
HEADER_FONT  = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
CELL_FONT    = Font(name="Calibri", size=10)
THIN_BORDER  = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)

ROLE_FILLS = {"user": USER_FILL, "assistant": AI_FILL}


def build_workbook(rows: list[dict]) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Chat History"

    # ── Header row ──────────────────────────────────────────────────────
    for col_idx, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font    = HEADER_FONT
        cell.fill    = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border  = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

    # ── Data rows ────────────────────────────────────────────────────────
    for row_idx, r in enumerate(rows, start=2):
        values = [
            r["user_id"],
            r["conv_id"],
            r["conv_title"],
            r["conv_created"],
            r["msg_num"],
            r["role"],
            r["timestamp"],
            r["content"],
        ]
        fill = ROLE_FILLS.get(r["role"], TOOL_FILL)
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font      = CELL_FONT
            cell.fill      = fill
            cell.border    = THIN_BORDER
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=(col_idx == len(COLUMNS)),  # wrap only content column
            )

    # ── Auto-filter ──────────────────────────────────────────────────────
    ws.auto_filter.ref = ws.dimensions

    return wb


async def fetch_history():
    # 1. Initialize Configuration
    if not app_settings.chat_history:
        print("Chat history is not configured in settings.")
        return

    print(f"Connecting to Cosmos DB: {app_settings.chat_history.account}")

    # 2. Authenticate
    if app_settings.chat_history.account_key:
        credential = app_settings.chat_history.account_key
    else:
        credential = DefaultAzureCredential()

    # 3. Initialize Client
    cosmos_endpoint = f"https://{app_settings.chat_history.account}.documents.azure.com:443/"
    client = CosmosConversationClient(
        cosmosdb_endpoint=cosmos_endpoint,
        credential=credential,
        database_name=app_settings.chat_history.database,
        container_name=app_settings.chat_history.conversations_container
    )

    success, msg = await client.ensure()
    if not success:
        print(f"Failed to connect: {msg}")
        return

    print("Connected successfully.")

    # 4. Fetch All Users
    print("Fetching all users...")
    query = "SELECT DISTINCT VALUE c.userId FROM c WHERE c.type = 'conversation'"
    users = []
    try:
        async for item in client.container_client.query_items(query=query):
            users.append(item)
    except Exception as e:
        print(f"Error fetching users: {e}")
        return

    print(f"Found {len(users)} users.")

    # 5. Collect all rows
    all_rows = []
    for user_id in users:
        print(f"  Fetching conversations for user: {user_id}")
        conversations = await client.get_conversations(user_id, limit=None)
        if not conversations:
            continue

        for conv in conversations:
            conv_id      = conv["id"]
            conv_title   = conv.get("title", "")
            conv_created = conv.get("createdAt", "")

            messages = await client.get_messages(user_id, conv_id)
            if not messages:
                # Still record the conversation with no messages
                all_rows.append({
                    "user_id":      user_id,
                    "conv_id":      conv_id,
                    "conv_title":   conv_title,
                    "conv_created": conv_created,
                    "msg_num":      "",
                    "role":         "",
                    "timestamp":    "",
                    "content":      "(no messages)",
                })
                continue

            for msg_num, msg in enumerate(messages, start=1):
                all_rows.append({
                    "user_id":      user_id,
                    "conv_id":      conv_id,
                    "conv_title":   conv_title,
                    "conv_created": conv_created,
                    "msg_num":      msg_num,
                    "role":         msg.get("role", ""),
                    "timestamp":    msg.get("createdAt", msg.get("updatedAt", "")),
                    "content":      msg.get("content", ""),
                })

    # 6. Write Excel
    timestamp  = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir    = os.path.join(os.path.dirname(__file__), '..', 'tools')
    out_path   = os.path.abspath(os.path.join(out_dir, f"chat_history_{timestamp}.xlsx"))

    wb = build_workbook(all_rows)
    wb.save(out_path)
    print(f"\nExported {len(all_rows)} rows to: {out_path}")

    # 7. Clean up
    if not app_settings.chat_history.account_key and hasattr(credential, 'close'):
        await credential.close()
    await client.cosmosdb_client.close()


if __name__ == "__main__":
    try:
        asyncio.run(fetch_history())
    except KeyboardInterrupt:
        pass
    except Exception as e:
        print(f"An error occurred: {e}")
